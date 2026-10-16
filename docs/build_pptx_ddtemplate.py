"""Build the Acme demo deck on the official Datadog presentation template (editable PPTX).

Usage: ~/.cache/slides-venv/bin/python docs/build_pptx_ddtemplate.py
Output: docs/otel-demo-slides-datadog-template.pptx

Template slides are reused in place (text replaced, unused shapes removed); slide 4 is duplicated for the
Collector slide. All other template slides are dropped.
"""
import copy
import os
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.text.text import _Run
from pptx.util import Inches, Pt

TEMPLATE = Path(os.environ["DD_PPTX_TEMPLATE"])  # path to Datadog_Presentation_Template.pptx
HERE = Path(__file__).parent
OUT = HERE / "otel-demo-slides-datadog-template.pptx"

PURPLE, VIOLET, TEAL, PINK, ORANGE = "632CA6", "5C45D5", "00B2A9", "E91E8C", "FF6B35"
INK, MUTED, WHITE = "1A1A2E", "5B5B73", "FFFFFF"
CODE_BG, CODE_FG, C_STR, C_COM, C_KEY, C_HL = "1D1133", "E6DAFF", "7FF0E6", "9AA0AA", "FF8F6B", "FFE27A"
MONO = "Consolas"

prs = Presentation(TEMPLATE)
T = {i + 1: s for i, s in enumerate(prs.slides)}  # template slide number -> slide


# ───────────────────────── helpers ─────────────────────────
def rgb(h):
    return RGBColor.from_string(h)


def shape(slide, shape_id):
    for sh in slide.shapes:
        if sh.shape_id == shape_id:
            return sh
    raise KeyError(shape_id)


def remove(slide, *ids):
    for i in ids:
        el = shape(slide, i)._element
        el.getparent().remove(el)


def style_run(r_el, o):
    f = _Run(r_el, None).font
    if "size" in o: f.size = Pt(o["size"])
    if "bold" in o: f.bold = o["bold"]
    if "italic" in o: f.italic = o["italic"]
    if "color" in o: f.color.rgb = rgb(o["color"])
    if "font" in o: f.name = o["font"]


def set_text(sh, paras, uniform=False, **defaults):
    """Replace a shape's text, keeping the template's paragraph/run formatting.

    paras: str | list of paragraphs; a paragraph is str or list of (text, opts) runs.
    defaults (size/bold/color/font/italic) apply to every run unless overridden per run.
    """
    tx = sh.text_frame._txBody
    old = tx.findall(qn("a:p"))
    fmts = []
    for p in old:
        r = p.find(qn("a:r"))
        rpr = r.find(qn("a:rPr")) if r is not None else p.find(qn("a:endParaRPr"))
        fmts.append((p.find(qn("a:pPr")), rpr))
    for p in old:
        tx.remove(p)
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        ppr, rpr = fmts[0 if uniform else min(i, len(fmts) - 1)]
        p = OxmlElement("a:p")
        tx.append(p)
        if ppr is not None:
            p.append(parse_xml(etree.tostring(ppr)))
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = OxmlElement("a:r")
            if rpr is not None:
                rp = copy.deepcopy(rpr)
                rp.tag = qn("a:rPr")
                r.append(parse_xml(etree.tostring(rp)))
            else:
                r.append(OxmlElement("a:rPr"))
            r.append(OxmlElement("a:t"))
            r[-1].text = t
            p.append(r)
            style_run(r, {**defaults, **o})
    return sh


def textbox(slide, x, y, w, h, paras, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT, font=None,
            anchor=MSO_ANCHOR.TOP, spacing=1.1, space_after=0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment, p.line_spacing, p.space_after = align, spacing, Pt(space_after)
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t
            base = {"size": size, "bold": bold, "color": color}
            if font:
                base["font"] = font
            style_run(r._r, {**base, **o})
    return tb


def rect(slide, x, y, w, h, fill=None, line=None, radius=0.06, kind=MSO_SHAPE.ROUNDED_RECTANGLE, line_w=1.25):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line); s.line.width = Pt(line_w)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def code(slide, x, y, w, h, lines, size=12, label=None, label_color=MUTED):
    if label:
        textbox(slide, x, y - 0.32, w, 0.3, label, size=11, color=label_color, font=MONO)
    rect(slide, x, y, w, h, fill=CODE_BG, radius=0.04)
    paras = [[(ln, {})] if isinstance(ln, str) else [(t, {"color": c}) for t, c in ln] for ln in lines]
    textbox(slide, x + 0.15, y + 0.1, w - 0.3, h - 0.18, [p or [(" ", {})] for p in paras],
            size=size, color=CODE_FG, font=MONO, spacing=1.15)


def notes(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt


def duplicate(src):
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)
    rid_map = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype.endswith(("/slideLayout", "/notesSlide")):
            continue
        rid_map[rid] = (new.part.relate_to(rel.target_ref, rel.reltype, is_external=True) if rel.is_external
                        else new.part.relate_to(rel.target_part, rel.reltype))
    bg = src._element.cSld.find(qn("p:bg"))
    if bg is not None:
        new._element.cSld.insert(0, copy.deepcopy(bg))
    tree = new.shapes._spTree
    for el in src.shapes._spTree.iterchildren():
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        tree.append(copy.deepcopy(el))
    for el in tree.iter():
        for attr in (qn("r:embed"), qn("r:link"), qn("r:id")):
            if el.get(attr) in rid_map:
                el.set(attr, rid_map[el.get(attr)])
    return new


S8 = duplicate(T[4])  # Collector slide: copy template slide 4 before slide 6 modifies it

# ───────────────────────── 1. Cover (template 1) ─────────────────────────
s = T[1]
set_text(shape(s, 1410), "One request. One trace.")
set_text(shape(s, 1409), ["Acme × Datadog · OpenTelemetry end to end",
                          "Angular · Java · Linux  —  live demo, October 2026"])
notes(s, "Cover. One request, one trace: Datadog RUM on Angular linked to OpenTelemetry-instrumented Java services, "
         "shipped via the OpenTelemetry Collector on Linux.")

# ───────────────────────── 2. Agenda (template 6) ─────────────────────────
s = T[6]
agenda = ["The challenge: requests hop A → B → C", "The approach: open standards in, one view out",
          "Angular: Datadog RUM, speaking W3C", "Java: OpenTelemetry agent, zero code",
          "Linux: OpenTelemetry Collector", "Live demo, what it means for Acme, next step"]
for gid, item in zip([1461, 1489, 1485, 1481, 1477, 1473], agenda):
    for sub in shape(s, gid).shapes:
        if sub.has_text_frame and "Body" in sub.text_frame.text:
            set_text(sub, item)
remove(s, 1469, 1465)  # items 07, 08
notes(s, "Agenda: challenge recap, approach, the three instrumentation layers, then the live demo and next step.")

# ───────────────────────── 3. Challenge (template 15) ─────────────────────────
s = T[15]
set_text(shape(s, 1623), "Requests hop A → B → C. No team sees past its own system.", size=30)
set_text(shape(s, 1625), "What we heard: a click travels Service A → B → C, and each team only has its own logs and metrics.",
         size=14)
pains = [(1630, "Where did the 3 seconds go?", "Each dashboard looks healthy alone",
          "Latency is argued, not measured"),
         (1627, "Whose error is it?", "A failure in C surfaces as a 500 in A",
          "Incidents bounce between teams"),
         (1633, "What did the customer see?", "Backend logs never meet the browser session",
          "No customer-to-code link")]
for sid, q, line, today in pains:
    set_text(shape(s, sid), [[(q, {"bold": True, "size": 17})], [(line, {"size": 13, "bold": False})]])
    sh = shape(s, sid)
    textbox(s, sh.left / 914400 + 0.05, 5.8, sh.width / 914400 - 0.1, 0.75,
            [[("TODAY", {"bold": True, "color": PURPLE, "size": 11})], [(today, {"bold": True})]],
            size=13, align=PP_ALIGN.CENTER)
notes(s, "Recap Acme's challenge: requests hop A → B → C and each team only sees its own system. "
         "The three pain points are examples — adjust to what Acme actually said.")

# ───────────────────────── 4. Approach (template 54) ─────────────────────────
s = T[54]
set_text(shape(s, 2406), "Open standards in, one view out")
steps = [(2412, "Browser", "Angular + Datadog RUM SDK"), (2413, "adt-service-a", "Spring Boot · jek-my1"),
         (2414, "service-b", "Spring Boot · jek-my2"), (2415, "service-c", "Spring Boot · jek-cz1"),
         (2416, "Datadog", "RUM · APM · Logs · Infra")]
for sid, main, sub in steps:
    set_text(shape(s, sid), [[(main, {"bold": True, "size": 15})], [(sub, {"size": 11})]])
for x in (2.76, 5.27, 7.77):
    textbox(s, x - 0.45, 2.45, 1.2, 0.3, "traceparent", size=10, color=PURPLE, bold=True, align=PP_ALIGN.CENTER)
textbox(s, 0.76, 1.45, 11.8, 0.6, "Instrument once with OpenTelemetry, keep vendor-neutral pipelines, and let Datadog "
        "stitch browser, services, hosts and logs together.", size=16, color=MUTED)
rect(s, 0.76, 5.0, 11.8, 1.75, fill="F3EFFE", radius=0.08)
textbox(s, 1.0, 5.1, 11.3, 1.55, [
    [("Every hop ", {}), ("carries the W3C traceparent header, ", {"bold": True}),
     ("injected by the RUM SDK in the browser and continued by the OpenTelemetry Java agent.", {})],
    [("Every host ", {}), ("runs an OpenTelemetry Collector ", {"bold": True}),
     ("that sends OTLP traces, metrics and logs plus host metrics and syslog to Datadog.", {})],
], size=14, spacing=1.15, space_after=4)
notes(s, "Architecture: browser (Angular + RUM SDK) → adt-service-a (jek-my1) → service-b (jek-my2) → service-c "
         "(jek-cz1). Each host runs otelcol-contrib, which sends OTLP traces, metrics and logs to Datadog.")

# ───────────────────────── 5. Divider (template 7) ─────────────────────────
s = T[7]
set_text(shape(s, 1497), "OpenTelemetry instrumentation", size=52)
set_text(shape(s, 1498), "Angular · Java · Linux — open standards, no code changes in the services")
notes(s, "Transition into the three instrumentation layers.")


# ───────────────────── 6–8. Instrumentation slides (templates 4, 5, 4) ─────────────────────
def panel(slide, title_id, body_id, title, bullets, lines_ids, pic_id, title_size=24):
    set_text(shape(slide, title_id), title, size=title_size)
    body = shape(slide, body_id)
    body.top, body.height = Inches(2.45), Inches(4.3)
    set_text(body, [[("•  ", {}), (b, {})] for b in bullets], uniform=True, size=14)
    for p in body.text_frame.paragraphs:
        p.space_after = Pt(10)
    remove(slide, *lines_ids, pic_id)


S_, C_, K_, H_ = C_STR, C_COM, C_KEY, C_HL
s = T[4]
panel(s, 1437, 1439, "Angular: Datadog RUM, speaking W3C",
      ["Every /api/* call carries traceparent, understood by any OpenTelemetry backend",
       "Views, actions, errors and Session Replay per user",
       "No change to Angular HttpClient calls"], (1440, 1441), 1442)
code(s, 4.89, 1.2, 8.0, 3.55, [
    [("datadogRum.", CODE_FG), ("init", K_), ("({", CODE_FG)],
    [("  applicationId: ", CODE_FG), ("'<RUM_APPLICATION_ID>'", S_), (", clientToken: ", CODE_FG), ("'<RUM_CLIENT_TOKEN>'", S_), (",", CODE_FG)],
    [("  site: ", CODE_FG), ("'datadoghq.com'", S_), (", service: ", CODE_FG), ("'shipping-express'", S_),
     (", env: ", CODE_FG), ("'uat'", S_), (",", CODE_FG)],
    [("  sessionSampleRate: 100, sessionReplaySampleRate: 100,", CODE_FG)],
    [("  trackResources: true, trackUserInteractions: true,", CODE_FG)],
    [("  allowedTracingUrls: [{", H_)],
    [("    match: (url) => url.startsWith(`${origin}/api/`),", H_)],
    [("    propagatorTypes: ['tracecontext'],   ", H_), ("// W3C traceparent", C_)],
    [("  }],", H_)],
    [("  plugins: [angularPlugin({ router: ", CODE_FG), ("true", K_), (" })],", CODE_FG)],
    [("});", CODE_FG)],
], size=12.5, label="apps/web/src/main.ts")
code(s, 4.89, 5.15, 8.0, 1.45, [
    [("providers: [ provideRouter(routes),", CODE_FG)],
    [("  provideDatadogRouter(),", H_), ("        // route change → RUM view", C_)],
    [("  provideDatadogErrorHandler(),", H_), ("  // Angular errors → RUM", C_)],
    [("  provideHttpClient(withFetch()) ]", CODE_FG)],
], size=12.5, label="apps/web/src/app/app.config.ts")
notes(s, "main.ts: allowedTracingUrls with propagatorTypes ['tracecontext'] makes the RUM SDK inject the W3C "
         "traceparent header into same-origin /api calls. provideDatadogRouter() is required with "
         "angularPlugin({router:true}) to record route changes as RUM views.")

s = T[5]
panel(s, 1448, 1450, "Java: OpenTelemetry agent, zero code",
      ["Agent v2.32.0 traces Tomcat, Spring MVC and the HTTP client automatically",
       "Continues the browser's W3C context and passes it on to B and C",
       "JVM metrics and trace-linked logs over the same OTLP pipe"], (1451, 1452), 1453)
shape(s, 1448).top = Inches(0.88)
code(s, 4.89, 1.2, 8.0, 3.15, [
    [("# systemd unit — who am I, and where do I send telemetry?", C_)],
    [("OTEL_SERVICE_NAME=", CODE_FG), ("adt-service-a", S_)],
    [("OTEL_EXPORTER_OTLP_ENDPOINT=", CODE_FG), ("http://localhost:4318", S_), ("   # local Collector", C_)],
    [("OTEL_RESOURCE_ATTRIBUTES=", CODE_FG), ("deployment.environment.name=uat,service.version=1.0.0", S_)],
    [("OTEL_TRACES_EXPORTER=", CODE_FG), ("otlp", S_), ("  OTEL_METRICS_EXPORTER=", CODE_FG), ("otlp", S_),
     ("  OTEL_LOGS_EXPORTER=", CODE_FG), ("otlp", S_)],
    [("OTEL_INSTRUMENTATION_COMMON_PEER_SERVICE_MAPPING=", CODE_FG), ("10.0.1.176=service-b", S_)],
    "",
    [("java -javaagent:/opt/app/opentelemetry-javaagent.jar -jar /opt/app/app.jar", H_)],
], size=12, label="apps/deploy/install-app.sh → systemd unit")
code(s, 4.89, 4.75, 8.0, 1.85, [
    [("@PostMapping(", CODE_FG), ("\"/shipments\"", S_), (")", CODE_FG)],
    [("public Map<String,Object> createShipment(@RequestBody Map<String,String> req) {", CODE_FG)],
    [("  log.info(", CODE_FG), ("\"Creating shipment from {} to {}\"", S_), (", …);", CODE_FG)],
    [("  serviceB.get().retrieve().body(…);  ", CODE_FG), ("// → service-b → service-c", C_)],
    [("  return Map.of(", CODE_FG), ("\"shipmentId\"", S_), (", id, ", CODE_FG), ("\"status\"", S_), (", ", CODE_FG),
     ("\"CREATED\"", S_), (");  } ", CODE_FG), ("// no OpenTelemetry imports", C_)],
], size=12, label="ShippingController.java — application code, unchanged")
notes(s, "The Java services have no OpenTelemetry code. One -javaagent flag plus environment variables give traces, "
         "JVM metrics and trace-linked logs, exported over OTLP to the local Collector.")

s = S8
panel(s, 1437, 1439, "Linux: OpenTelemetry Collector",
      ["otelcol-contrib 0.149.0 as a systemd service on Ubuntu 22.04",
       "Replaces a vendor agent for host metrics, syslog and app telemetry",
       "API key in a root-only env file, never in YAML"], (1440, 1441), 1442)
stages = [("Receivers", "otlp · hostmetrics · filelog", PURPLE), ("Processors", "resourcedetection · resource", TEAL),
          ("Connector", "datadog/connector → APM stats", VIOLET), ("Exporter", "datadog/exporter + extension", PINK)]
sw, gap = 1.82, 0.24
for i, (t, d, col) in enumerate(stages):
    x = 4.89 + i * (sw + gap)
    rect(s, x, 1.2, sw, 1.45, fill=WHITE, line=col, radius=0.08, line_w=2)
    textbox(s, x + 0.08, 1.32, sw - 0.16, 0.4, [[(t, {"bold": True})]], size=15, color=col, align=PP_ALIGN.CENTER)
    textbox(s, x + 0.08, 1.75, sw - 0.16, 0.8, d, size=11.5, color=INK, align=PP_ALIGN.CENTER)
    if i < 3:
        textbox(s, x + sw - 0.02, 1.62, gap + 0.04, 0.5, [[("→", {"bold": True})]], size=18, color=PURPLE,
                align=PP_ALIGN.CENTER)
code(s, 4.89, 3.25, 8.0, 2.05, [
    [("service:", CODE_FG)],
    [("  extensions: [health_check, file_storage, ", CODE_FG), ("datadog", H_), ("]   ", CODE_FG), ("# Fleet Automation", C_)],
    [("  pipelines:", CODE_FG)],
    [("    traces:   ", CODE_FG), ("otlp", S_), (" → resourcedetection, resource → ", CODE_FG), ("datadog/connector", H_)],
    [("    traces/2: ", CODE_FG), ("datadog/connector", H_), (" → ", CODE_FG), ("datadog/exporter", S_)],
    [("    metrics:  ", CODE_FG), ("otlp, hostmetrics, datadog/connector", S_), (" → datadog/exporter", CODE_FG)],
    [("    logs:     ", CODE_FG), ("otlp, filelog/system", S_), (" → datadog/exporter", CODE_FG)],
], size=12, label="otelcol-ubuntu/config.yaml")
rect(s, 4.89, 5.55, 8.0, 1.05, fill="F3EFFE", radius=0.08)
textbox(s, 5.1, 5.62, 7.6, 0.9, [[("Vendor-neutral by design: ", {"bold": True}),
        ("receivers and processors are portable OpenTelemetry components. See the full pipeline live in OTelBin "
         "(demo step 6).", {})]], size=14, anchor=MSO_ANCHOR.MIDDLE, spacing=1.15)
notes(s, "Walk the pipeline: OTLP + hostmetrics + filelog receivers → resourcedetection/resource processors → "
         "datadog/connector computes APM stats → datadog/exporter. The datadog extension shows the Collector in "
         "Fleet Automation.")

# ───────────────────────── 9. How it stitches (template 27) ─────────────────────────
s = T[27]
set_text(shape(s, 1859), "One click → one trace across three teams")
body = shape(s, 1860)
body.height = Inches(0.6)
set_text(body, "Real trace from the demo (14911808522202188851): a customer taps Track, and service-c's injected "
               "failure is visible end to end.", size=16, bold=False)
code(s, 0.76, 2.55, 11.8, 0.8, [
    [("// injected by the RUM SDK on the browser request", C_)],
    [("traceparent: 00-", CODE_FG), ("0000000000000000cef162d844810c33", H_), ("-1d96e32a32448b12-01", CODE_FG),
     ("      tracestate: dd=s:1;o:rum", CODE_FG)],
], size=13)
rows = [("shipping-express · browser", 0, 100, "F5B700", "GET /api/track · 409 ms · 500"),
        ("  └ adt-service-a · jek-my1", 67.3, 26.2, PURPLE, "107 ms"),
        ("      └ service-b · jek-my2", 71.7, 17.5, TEAL, "72 ms"),
        ("          └ service-c · jek-cz1", 73.5, 13.8, PINK, "57 ms · error")]
tx, tw = 4.6, 12.56 - 4.6
for i, (lab, left, width, col, txt) in enumerate(rows):
    y = 3.6 + i * 0.5
    textbox(s, 0.76, y, 3.8, 0.42, lab, size=12.5, font=MONO, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, tx, y, tw, 0.42, fill="EEEAF7", radius=0.15)
    rect(s, tx + tw * left / 100, y + 0.04, tw * width / 100, 0.34, fill=col, radius=0.2)
    bar_x, bar_w = tx + tw * left / 100, tw * width / 100
    label_w = len(txt) * 11.5 * 0.0085 + 0.2
    if label_w < bar_w:                                  # inside the bar
        lx, lw, align, color = bar_x + 0.08, bar_w, PP_ALIGN.LEFT, INK if col == "F5B700" else WHITE
    elif bar_x + bar_w + 0.08 + label_w <= tx + tw:      # right of the bar
        lx, lw, align, color = bar_x + bar_w + 0.08, label_w, PP_ALIGN.LEFT, INK
    else:                                                # left of the bar
        lx, lw, align, color = bar_x - label_w - 0.08, label_w, PP_ALIGN.RIGHT, INK
    textbox(s, lx, y, lw, 0.42, [[(txt, {"bold": True})]], size=11.5, color=color, align=align,
            anchor=MSO_ANCHOR.MIDDLE)
cw = (11.8 - 0.4) / 3
for i, (q, a) in enumerate([("Where did the time go?", "107 ms in the services; ~300 ms was browser queueing and network."),
                            ("Whose error is it?", "service-c on jek-cz1, with the stack trace attached."),
                            ("What did the customer see?", "Their Session Replay, at the exact moment.")]):
    x = 0.76 + i * (cw + 0.2)
    rect(s, x, 5.75, cw, 1.05, fill="F3EFFE", radius=0.1)
    textbox(s, x + 0.18, 5.83, cw - 0.36, 0.9, [[(q, {"bold": True, "color": PURPLE})], [(a, {})]], size=13, spacing=1.12)
notes(s, "Real trace 14911808522202188851. The RUM SDK injected traceparent; adt-service-a's span is the direct child "
         "of the browser request, and service-c's simulated failure propagates back up. Bar positions use the browser's "
         "real first-byte timing; backend offsets are approximate.")

# ───────────────────────── 10. Live demo (template 34) ─────────────────────────
s = T[34]
set_text(shape(s, 1960), "Live demo: let's follow a shipment")
remove(s, 1962, 1963)
html = (HERE / "otel-demo-slides.html").read_text()
otelbin = html[html.index('href="https://www.otelbin.io/') + 6:]
otelbin = otelbin[:otelbin.index('"')]
demos = [
    ("The customer's view", "Create a shipment, get a quote, track a parcel", "<MY1_PUBLIC_IP>:8080",
     "http://<MY1_PUBLIC_IP>:8080/"),
    ("Session Replay", "Watch exactly what the customer did and saw", "RUM → Session Replay",
     "https://app.datadoghq.com/rum/sessions"),
    ("Click → backend trace", "From the RUM request into the A → B → C waterfall", "RUM → Trace waterfall",
     "https://app.datadoghq.com/rum/sessions"),
    ("Service map", "Dependencies discovered automatically from OTel traces", "Software Catalog → Map",
     "https://app.datadoghq.com/software"),
    ("Shared dashboard", "Requests, errors, latency, dependencies, hosts, logs", "Dashboards → ADT Demo",
     "https://app.datadoghq.com/dashboard/lists"),
    ("The Collector config", "The real pipeline, visualised", "OTelBin", otelbin),
]
frame = shape(s, 1959)
tbl = s.shapes.add_table(len(demos) + 1, 4, frame.left, frame.top, frame.width, Inches(5.49)).table
for col, w in enumerate([0.75, 3.1, 4.9, 3.04]):
    tbl.columns[col].width = Inches(w)
tbl.first_row = True
for c, head in enumerate(["#", "Step", "What you'll show", "Open"]):
    cell = tbl.cell(0, c)
    cell.fill.solid(); cell.fill.fore_color.rgb = rgb(PURPLE)
    cell.text_frame.text = head
    style_run(cell.text_frame.paragraphs[0].runs[0]._r, {"bold": True, "size": 14, "color": WHITE})
for r_i, (step, what, label, href) in enumerate(demos, start=1):
    vals = [str(r_i), step, what, label]
    for c, v in enumerate(vals):
        cell = tbl.cell(r_i, c)
        cell.fill.solid(); cell.fill.fore_color.rgb = rgb("F3EFFE" if r_i % 2 == 0 else WHITE)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        run = cell.text_frame.paragraphs[0].add_run()
        run.text = v
        style_run(run._r, {"size": 14, "color": PURPLE if c in (0, 3) else INK, "bold": c in (0, 1)})
        if c == 3:
            run.hyperlink.address = href
tbl.rows[0].height = Inches(0.6)
for r_i in range(1, len(demos) + 1):
    tbl.rows[r_i].height = Inches(0.8)
notes(s, "Demo order — the Open column links are clickable:\n" + "\n".join(f"{i + 1}. {d[0]}: {d[3]}" for i, d in enumerate(demos)))

# ───────────────────────── 11. What it means (template 39) ─────────────────────────
s = T[39]
set_text(shape(s, 2045), "What this means for Acme")
set_text(shape(s, 2046), " ")
recap = [(2052, 2056, ("Open by", "default"), ["OpenTelemetry agent and Collector, no proprietary agent on the hosts",
                                                "W3C Trace Context end to end, portable config"]),
         (2053, 2058, ("Faster", "triage"), ["One trace from browser click to the failing hop",
                                              "Logs, host metrics and errors on the same trace ID"]),
         (2054, 2060, ("Customer", "first"), ["RUM and Session Replay show what customers experienced",
                                               "Errors linked to the backend request that caused them"])]
remove(s, 2059)  # template "$" icon doesn't fit "Faster triage"
shape(s, 2058).top = Inches(3.45)
for body_id, label_id, label, items in recap:
    set_text(shape(s, body_id), items)
    set_text(shape(s, label_id), list(label))
notes(s, "Close: open by default, faster triage, customer-first.")

# ───────────────────────── 12. Next step (template 41) ─────────────────────────
s = T[41]
set_text(shape(s, 2078), "Next step: a 2-week PoC")
set_text(shape(s, 2081), "Pick one real A → B → C shipment flow", size=26)
set_text(shape(s, 2080), "Instrument it with RUM · OTel agent · Collector", size=26)
set_text(shape(s, 2079), "Success: one trace from click to failing hop", size=26)
notes(s, "Ask: a 2-week PoC on one real A → B → C flow, using exactly this instrumentation pattern.")

# ───────────────────────── order + drop unused template slides ─────────────────────────
order = [T[1], T[6], T[15], T[54], T[7], T[4], T[5], S8, T[27], T[34], T[39], T[41]]
lst = prs.slides._sldIdLst
by_part = {prs.part.related_part(sid.rId): sid for sid in lst}
keep = [by_part[sl.part] for sl in order]
for sid in list(lst):
    if sid not in keep:
        prs.part.drop_rel(sid.rId)
    lst.remove(sid)
for sid in keep:
    lst.append(sid)

prs.save(OUT)
print("saved", OUT, f"({OUT.stat().st_size // 1024} KB, {len(prs.slides)} slides)")
