"""Build an editable PPTX version of docs/otel-demo-slides.html with python-pptx.

Usage: ~/.cache/slides-venv/bin/python docs/build_pptx.py
Output: docs/otel-demo-slides-editable.pptx
"""
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

OUT = Path(__file__).with_name("otel-demo-slides-editable.pptx")

# Datadog palette (+ Acme accents)
BG = "0D1117"
PURPLE, PURPLE_LT, PURPLE_TINT = "632CA6", "7B3FC4", "D0A8FF"
ORANGE, TEAL, BLUE = "FF6B35", "00C7B7", "3E9BFF"
YELLOW, RED = "F5B700", "E0533D"
WHITE, MUTED, GREY = "FFFFFF", "A3A7AE", "7D8590"
CODE_BG, CODE_FG = "0D0920", "E6DAFF"
FONT, MONO = "Calibri", "Consolas"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
L, R = 0.6, 13.333 - 0.6  # content margins (in)
W = R - L


def rgb(hex_):
    return RGBColor.from_string(hex_)


def set_alpha(fill_parent_xml, alpha_pct):
    """Apply transparency to the first solidFill colour of a shape fill/line."""
    clr = fill_parent_xml.find(qn("a:solidFill"))[0]
    for old in clr.findall(qn("a:alpha")):
        clr.remove(old)
    etree.SubElement(clr, qn("a:alpha")).set("val", str(int(alpha_pct * 1000)))


def box(slide, x, y, w, h, fill=None, fill_alpha=100, line=None, line_alpha=100, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
        if fill_alpha < 100:
            set_alpha(shp.fill._xPr, fill_alpha)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(1.25)
        if line_alpha < 100:
            set_alpha(shp.line._get_or_add_ln(), line_alpha)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def text(slide, x, y, w, h, paras, size=16, color=WHITE, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=FONT, spacing=1.1, space_after=0):
    """paras: str | list of paragraphs; each paragraph is str or list of (text, {opts}) runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(space_after)
        runs = [(para, {})] if isinstance(para, str) else para
        for t, o in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = rgb(o.get("color", color))
            if o.get("link"):
                r.hyperlink.address = o["link"]
    return tb


def bullets(slide, x, y, w, h, items, size=15, color="DADCE0"):
    paras = [[("•  ", {"color": ORANGE, "bold": True}), *([(it, {})] if isinstance(it, str) else it)] for it in items]
    return text(slide, x, y, w, h, paras, size=size, color=color, spacing=1.12, space_after=7)


def badge_w(label, size=11):
    return max(0.55, 0.26 + len(label) * size * 0.0082)


def badge(slide, x, y, label, kind="p", size=11):
    colors = {"p": (PURPLE, PURPLE_LT, PURPLE_TINT), "o": (ORANGE, ORANGE, "FFB596"),
              "g": (TEAL, TEAL, "7FF0E6"), "y": (YELLOW, YELLOW, "FFE27A")}
    fill, line, fg = colors[kind]
    w = badge_w(label, size)
    b = box(slide, x, y, w, 0.3, fill=fill, fill_alpha=30, line=line, line_alpha=60, radius=0.5)
    tf = b.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size, r.font.bold, r.font.name = Pt(size), True, FONT
    r.font.color.rgb = rgb(fg)
    return w


def badges_row(slide, x, y, labels, kind="p", size=11, gap=0.1, max_x=None):
    cx = x
    for lab in labels:
        w = badge_w(lab, size)
        if max_x and cx + w > max_x:
            cx, y = x, y + 0.38
        badge(slide, cx, y, lab, kind, size)
        cx += w + gap


def code(slide, x, y, w, h, lines, size=12.5, label=None):
    if label:
        text(slide, x, y - 0.34, w, 0.3, label, size=11, color=MUTED, font=MONO)
    box(slide, x, y, w, h, fill=CODE_BG, line="FFFFFF", line_alpha=12, radius=0.04)
    paras = []
    for ln in lines:
        runs = [(ln, {})] if isinstance(ln, str) else [(t, {"color": c}) for t, c in ln]
        paras.append(runs or [(" ", {})])
    text(slide, x + 0.15, y + 0.12, w - 0.3, h - 0.2, paras, size=size, color=CODE_FG, font=MONO, spacing=1.18)


def card(slide, x, y, w, h, title=None, body=None, outcome=None, hl=False, title_size=19, body_size=14.5, border=None,
         title_lines=None):
    if hl:
        box(slide, x, y, w, h, fill=PURPLE, fill_alpha=22, line=PURPLE_LT, line_alpha=55)
    else:
        box(slide, x, y, w, h, fill="FFFFFF", fill_alpha=6, line=border or "FFFFFF", line_alpha=100 if border else 14)
    cy = y + 0.2
    if title:
        text(slide, x + 0.22, cy, w - 0.44, 0.75, [[(title, {"bold": True})]], size=title_size)
        wraps = title_lines == 2 if title_lines else len(title) * title_size * 0.0066 > (w - 0.44)
        cy += 0.5 + (0.32 if wraps else 0)
    if body:
        if isinstance(body, list):
            bullets(slide, x + 0.22, cy, w - 0.44, h - (cy - y) - 0.95, body, size=body_size)
        else:
            text(slide, x + 0.22, cy, w - 0.44, h - (cy - y) - 0.95, body, size=body_size, color="DADCE0", spacing=1.15)
    if outcome:
        k, v = outcome
        oy = y + h - 0.88
        box(slide, x + 0.18, oy, w - 0.36, 0.7, fill="FFFFFF", fill_alpha=8, radius=0.12)
        text(slide, x + 0.3, oy + 0.06, w - 0.6, 0.25, [[(k.upper(), {"bold": True})]], size=9.5, color=GREY)
        text(slide, x + 0.3, oy + 0.3, w - 0.6, 0.36, [[(v, {"bold": True})]], size=14)


def node(slide, x, y, w, h, title, desc, border, tint=None, title_size=17):
    box(slide, x, y, w, h, fill=tint or "FFFFFF", fill_alpha=8 if tint else 4, line=border, radius=0.1)
    text(slide, x, y + 0.12, w, 0.4, [[(title, {"bold": True})]], size=title_size, align=PP_ALIGN.CENTER)
    text(slide, x + 0.08, y + 0.55, w - 0.16, h - 0.6, desc.split("\n"), size=12, color="C9CDD3", align=PP_ALIGN.CENTER, spacing=1.1)


def arrow_label(slide, x, y, w, label=None, color=ORANGE, glyph="→"):
    text(slide, x, y, w, 0.45, [[(glyph, {"bold": True})]], size=20, color=color, align=PP_ALIGN.CENTER)
    if label:
        text(slide, x - 0.15, y + 0.5, w + 0.3, 0.5, [[(ln, {"bold": True})] for ln in label.split("\n")],
             size=10.5, color=color, align=PP_ALIGN.CENTER, spacing=1.0)


def new_slide(num):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    # brand strip (purple → orange)
    box(s, 0, 0, 4.4, 0.06, fill=PURPLE_LT, shape=MSO_SHAPE.RECTANGLE)
    box(s, 4.4, 0, 4.5, 0.06, fill=PURPLE, shape=MSO_SHAPE.RECTANGLE)
    box(s, 8.9, 0, 4.433, 0.06, fill=ORANGE, shape=MSO_SHAPE.RECTANGLE)
    text(s, 12.0, 7.1, 0.9, 0.3, str(num), size=10, color="55556A", align=PP_ALIGN.RIGHT)
    return s


def header(s, eyebrow, title_runs, sub=None, y=0.45):
    text(s, L, y, W, 0.3, [[(eyebrow.upper(), {"bold": True})]], size=12, color=ORANGE)
    text(s, L, y + 0.32, W, 1.1, title_runs, size=34, bold=True, spacing=1.0)
    if sub:
        text(s, L, y + 1.35 if len(title_runs) > 1 else y + 0.98, W, 0.6, sub, size=15, color=MUTED, spacing=1.15)


def notes(s, txt):
    s.notes_slide.notes_text_frame.text = txt


# ───────────────────────── 1. Cover ─────────────────────────
s = new_slide(1)
box(s, L, 1.0, 0.95, 0.5, fill=YELLOW, radius=0.08)
text(s, L, 1.0, 0.95, 0.5, [[("Acme", {"bold": True, "italic": True, "color": RED})]], size=20, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, L + 1.05, 1.0, 0.4, 0.5, "×", size=20, color=MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
box(s, L + 1.5, 1.0, 1.55, 0.5, fill=PURPLE, radius=0.08)
text(s, L + 1.5, 1.0, 1.55, 0.5, [[("Datadog", {"bold": True})]], size=20, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, L, 1.9, W, 0.35, [[("LIVE DEMO · OPENTELEMETRY END TO END", {"bold": True})]], size=13, color=ORANGE)
text(s, L, 2.3, W, 1.8, [[("One request. One trace.", {"bold": True})],
                         [("From the click to the last hop.", {"bold": True, "color": "B98CFF"})]], size=46, spacing=1.0)
text(s, L, 4.1, 10.5, 0.8, "Datadog RUM on an Angular front end, linked to Java services instrumented with OpenTelemetry "
     "and shipped through the OpenTelemetry Collector on Linux.", size=16, color=MUTED, spacing=1.15)
cw = (W - 0.4) / 3
for i, (n, lab) in enumerate([("0", "lines of tracing code in the Java services"),
                              ("1", "W3C traceparent header links browser to backend"),
                              ("3 → 1", "three teams' systems, one shared view")]):
    x = L + i * (cw + 0.2)
    box(s, x, 5.2, cw, 1.5, fill="FFFFFF", fill_alpha=6, line="FFFFFF", line_alpha=14)
    text(s, x + 0.25, 5.3, cw - 0.5, 0.75, [[(n, {"bold": True})]], size=34, color="B98CFF")
    text(s, x + 0.25, 6.05, cw - 0.5, 0.6, lab, size=14, color="DADCE0")
notes(s, "Cover. One request, one trace: Datadog RUM on Angular linked to OpenTelemetry-instrumented Java services, "
         "shipped via the OpenTelemetry Collector on Linux.")

# ───────────────────────── 2. Challenge ─────────────────────────
s = new_slide(2)
header(s, "What we heard", [[("Requests hop ", {}), ("A → B → C", {"color": YELLOW}), (".", {})],
                            [("No team sees past its own system.", {})]])
ny, nh = 2.05, 1.05
cols = [(L, 1.55), None, (L + 2.25, 2.55), None, (L + 5.55, 2.55), None, (L + 8.85, 2.55)]
node(s, L, ny, 1.55, nh, "Customer", "clicks “Track”", "888888")
arrow_label(s, L + 1.6, ny + 0.2, 0.6)
node(s, L + 2.25, ny, 2.55, nh, "Service A", "Team A's logs & metrics", BLUE, BLUE)
arrow_label(s, L + 4.85, ny + 0.1, 0.65, "blind\nspot", RED, "║")
node(s, L + 5.55, ny, 2.55, nh, "Service B", "Team B's logs & metrics", TEAL, TEAL)
arrow_label(s, L + 8.15, ny + 0.1, 0.65, "blind\nspot", RED, "║")
node(s, L + 8.85, ny, W - 8.85, nh, "Service C", "Team C's logs & metrics", ORANGE, ORANGE)
cw, cy, ch = (W - 0.4) / 3, 3.45, 3.55
pains = [("“Where did the 3 seconds go?”", "Each team's dashboard looks healthy in isolation. Nobody can see the whole journey of one request.", "Latency is argued, not measured"),
         ("“Whose error is it?”", "A failure deep in C shows up as a generic 500 in A. Incidents bounce between teams via chat and screenshots.", "Hours of war-room triage"),
         ("“What did the customer see?”", "Backend logs never connect to the actual browser session, so a complaint can't be tied to a request.", "No customer-to-code link")]
for i, (t, b, o) in enumerate(pains):
    card(s, L + i * (cw + 0.2), cy, cw, ch, t, b, ("Today", o), title_lines=2)  # same offset across the row
notes(s, "Recap Acme's challenge: requests hop A → B → C, and each team only sees its own system. "
         "The three pain points are examples — adjust to what Acme actually said.")

# ───────────────────────── 3. Approach ─────────────────────────
s = new_slide(3)
header(s, "The approach", [[("Open standards in, one view out", {})]],
       "Instrument once with OpenTelemetry, keep vendor-neutral pipelines, and let Datadog stitch browser, services, hosts and logs together.")
aw, nw0 = 0.85, 2.35
nw = (W - nw0 - 3 * aw) / 3
xs = [L, L + nw0 + aw, L + nw0 + aw + nw + aw, L + nw0 + 2 * (aw + nw) + aw]
ty = 2.2
node(s, xs[0], ty, nw0, 1.15, "Browser", "Angular app\nDatadog RUM SDK", YELLOW, YELLOW)
for i, (name, desc, col) in enumerate([("adt-service-a", "Spring Boot · jek-my1\nOTel Java agent", BLUE),
                                       ("service-b", "Spring Boot · jek-my2\nOTel Java agent", TEAL),
                                       ("service-c", "Spring Boot · jek-cz1\nOTel Java agent", ORANGE)]):
    node(s, xs[i + 1], ty, nw, 1.15, name, desc, col)
    arrow_label(s, xs[i + 1] - aw, ty + 0.25, aw, "traceparent")
for i in range(4):
    cx, cwid = xs[i], (nw0 if i == 0 else nw)
    arrow_label(s, cx, ty + 1.2, cwid, "RUM events" if i == 0 else "OTLP", ORANGE, "↓")
    if i:
        node(s, cx, ty + 2.0, nw, 0.85, "OTel Collector", "hostmetrics · syslog", GREY, title_size=14)
box(s, L, 5.35, W, 1.55, fill=PURPLE, fill_alpha=22, line=PURPLE_LT, line_alpha=55)
text(s, L, 5.45, W, 0.5, [[("Datadog", {"bold": True})]], size=26, align=PP_ALIGN.CENTER)
text(s, L, 5.93, W, 0.35, "One trace ID across browser, services, hosts and logs", size=14, color="DADCE0", align=PP_ALIGN.CENTER)
labels = ["RUM", "Session Replay", "APM / Distributed Tracing", "Service Map", "Logs", "Infrastructure", "Fleet Automation", "Dashboards"]
total = sum(badge_w(l) for l in labels) + 0.1 * (len(labels) - 1)
badges_row(s, L + (W - total) / 2, 6.38, labels)
notes(s, "Architecture: browser (Angular + RUM SDK) → adt-service-a (jek-my1) → service-b (jek-my2) → service-c (jek-cz1). "
         "Each host runs otelcol-contrib, which sends OTLP traces, metrics and logs to Datadog.")

# ───────────────────────── 4. Angular ─────────────────────────
s = new_slide(4)
header(s, "Instrumentation 1 of 3 · Front end", [[("Angular: Datadog RUM, speaking W3C", {})]])
S_, C_, K_, H_ = "7FF0E6", GREY, "FF8F6B", "FFE27A"
code(s, L, 1.95, 7.3, 4.95, [
    [("datadogRum.", CODE_FG), ("init", K_), ("({", CODE_FG)],
    [("  applicationId: ", CODE_FG), ("'<RUM_APPLICATION_ID>'", S_), (",", CODE_FG)],
    [("  clientToken:   ", CODE_FG), ("'<RUM_CLIENT_TOKEN>'", S_), (",", CODE_FG)],
    [("  site: ", CODE_FG), ("'datadoghq.com'", S_), (",", CODE_FG)],
    [("  service: ", CODE_FG), ("'shipping-express'", S_), (", env: ", CODE_FG), ("'uat'", S_), (", version: ", CODE_FG), ("'1.0.0'", S_), (",", CODE_FG)],
    [("  sessionSampleRate: 100, sessionReplaySampleRate: 100,", CODE_FG)],
    [("  trackResources: true, trackUserInteractions: true,", CODE_FG)],
    [("  defaultPrivacyLevel: ", CODE_FG), ("'mask-user-input'", S_), (",", CODE_FG)],
    [("  allowedTracingUrls: [{", H_)],
    [("    match: (url) => url.startsWith(`${origin}/api/`),", H_)],
    [("    propagatorTypes: ['tracecontext'],   ", H_), ("// W3C traceparent", C_)],
    [("  }],", H_)],
    [("  plugins: [angularPlugin({ router: ", CODE_FG), ("true", K_), (" })],", CODE_FG)],
    [("});", CODE_FG)],
], size=13, label="apps/web/src/main.ts")
rx, rw = L + 7.55, W - 7.55
code(s, rx, 1.95, rw, 1.75, [
    [("providers: [", CODE_FG)],
    [("  provideRouter(routes),", CODE_FG)],
    [("  provideDatadogRouter(),", H_), ("  // route → RUM view", C_)],
    [("  provideDatadogErrorHandler(),", H_), (" // errors", C_)],
    [("  provideHttpClient(withFetch()),", CODE_FG)],
    [("]", CODE_FG)],
], size=11, label="apps/web/src/app/app.config.ts")
card(s, rx, 3.9, rw, 3.0, body=[
    [("Every ", {}), ("/api/*", {"font": MONO}), (" call carries ", {}), ("traceparent", {"font": MONO}), (", understood by OTel backends out of the box", {})],
    "Views, actions, errors and Session Replay per user",
    [("No change to Angular ", {}), ("HttpClient", {"font": MONO}), (" calls", {})],
], outcome=("Key outcome", "The trace starts in the customer's browser"))
notes(s, "Show main.ts: allowedTracingUrls with propagatorTypes ['tracecontext'] makes the RUM SDK inject the W3C traceparent "
         "header into same-origin /api calls. provideDatadogRouter() is required with angularPlugin({router:true}) "
         "to record route changes as RUM views.")

# ───────────────────────── 5. Java ─────────────────────────
s = new_slide(5)
header(s, "Instrumentation 2 of 3 · Java services", [[("Spring Boot: OpenTelemetry Java agent, zero code", {})]])
hw = (W - 0.3) / 2
code(s, L, 1.95, hw, 4.95, [
    [("# Who am I, and where do I send telemetry?", C_)],
    [("OTEL_SERVICE_NAME=", CODE_FG), ("adt-service-a", S_)],
    [("OTEL_EXPORTER_OTLP_ENDPOINT=", CODE_FG), ("http://localhost:4318", S_)],
    [("OTEL_RESOURCE_ATTRIBUTES=", CODE_FG), ("deployment.environment.name=uat,", S_)],
    [("                         ", CODE_FG), ("service.version=1.0.0", S_)],
    [("OTEL_TRACES_EXPORTER=", CODE_FG), ("otlp", S_)],
    [("OTEL_METRICS_EXPORTER=", CODE_FG), ("otlp", S_)],
    [("OTEL_LOGS_EXPORTER=", CODE_FG), ("otlp", S_), ("   # logs carry trace_id", C_)],
    [("# Name the downstream behind an IP", C_)],
    [("OTEL_INSTRUMENTATION_COMMON_PEER_SERVICE_MAPPING=", CODE_FG)],
    [("                         ", CODE_FG), ("10.0.1.176=service-b", S_)],
    "",
    [("java -javaagent:opentelemetry-javaagent.jar -jar app.jar", H_)],
], size=12.5, label="apps/deploy/install-app.sh → systemd unit")
rx = L + hw + 0.3
code(s, rx, 1.95, hw, 2.3, [
    [("@PostMapping(", CODE_FG), ("\"/shipments\"", S_), (")", CODE_FG)],
    [("public Map<String,Object> createShipment(", CODE_FG)],
    [("        @RequestBody Map<String,String> req) {", CODE_FG)],
    [("  log.info(", CODE_FG), ("\"Creating shipment from {} to {}\"", S_), (", …);", CODE_FG)],
    [("  serviceB.get().retrieve().body(…); ", CODE_FG), ("// → B → C", C_)],
    [("  return Map.of(", CODE_FG), ("\"shipmentId\"", S_), (", id, …);", CODE_FG)],
    [("}", CODE_FG)],
    [("// No OpenTelemetry imports. No spans. No headers.", C_)],
], size=12, label="ShippingController.java (application code, unchanged)")
card(s, rx, 4.4, hw, 2.5, body=[
    "Agent v2.32.0 traces Tomcat, Spring MVC and the HTTP client automatically",
    "Continues the browser's W3C context and passes it on to B and C",
    "JVM metrics and trace-linked logs over the same OTLP pipe",
], outcome=("Key outcome", "Teams change nothing in their code"), body_size=13.5)
notes(s, "The Java services have no OpenTelemetry code. One -javaagent flag plus environment variables give traces, "
         "JVM metrics and trace-linked logs, exported over OTLP to the local Collector.")

# ───────────────────────── 6. Collector ─────────────────────────
s = new_slide(6)
header(s, "Instrumentation 3 of 3 · Linux hosts", [[("OpenTelemetry Collector: one pipeline per host", {})]],
       "otelcol-contrib 0.149.0 as a systemd service on Ubuntu 22.04. It replaces a vendor agent for hosts, logs and app telemetry.")
aw = 0.45
sw = (W - 3 * aw) / 4
stages = [("Receivers", BLUE, "p", [("otlp", "apps · 4317 / 4318"), ("hostmetrics", "CPU · mem · disk · net"), ("filelog", "syslog · auth.log")]),
          ("Processors", TEAL, "g", [("resourcedetection", "host · OS"), ("resource", "host.name · env=uat")]),
          ("Connector", PURPLE_LT, "p", [("datadog/connector", "APM stats (requests, errors, latency) computed in the Collector")]),
          ("Exporter", ORANGE, "o", [("datadog/exporter", "traces · metrics · logs"), ("datadog extension", "Fleet Automation")])]
for i, (t, col, kind, items) in enumerate(stages):
    x = L + i * (sw + aw)
    box(s, x, 2.15, sw, 2.55, fill="FFFFFF", fill_alpha=6, line=col)
    text(s, x + 0.2, 2.25, sw - 0.4, 0.4, [[(t, {"bold": True})]], size=19)
    yy = 2.75
    for lab, desc in items:
        badge(s, x + 0.2, yy, lab, kind)
        text(s, x + 0.2, yy + 0.32, sw - 0.4, 0.55, desc, size=12.5, color="DADCE0")
        yy += 0.62 if len(desc) < 28 else 0.95
    if i < 3:
        arrow_label(s, x + sw, 3.15, aw)
hw = (W - 0.3) / 2
code(s, L, 5.3, hw * 1.12, 1.3, [
    [("traces:   ", CODE_FG), ("otlp", S_), (" → resourcedetection, resource → ", CODE_FG), ("datadog/connector", H_)],
    [("traces/2: ", CODE_FG), ("datadog/connector", H_), (" → ", CODE_FG), ("datadog/exporter", S_)],
    [("metrics:  ", CODE_FG), ("otlp, hostmetrics, datadog/connector", S_), (" → exporter", CODE_FG)],
    [("logs:     ", CODE_FG), ("otlp, filelog/system", S_), (" → exporter", CODE_FG)],
], size=12, label="otelcol-ubuntu/config.yaml · service.pipelines")
cx = L + hw * 1.12 + 0.3
card(s, cx, 4.85, R - cx, 2.05, body=[
    "Vendor-neutral config: receivers and processors are portable",
    "API key in a root-only env file, never in YAML",
    "Full config visualised live in OTelBin (demo step 6)",
], outcome=("Key outcome", "Hosts, logs and traces through one open pipeline"), body_size=12.5)
notes(s, "Walk the pipeline: OTLP + hostmetrics + filelog receivers → resourcedetection/resource processors → "
         "datadog/connector computes APM stats → datadog/exporter. The datadog extension shows the Collector in Fleet Automation. "
         "Open the OTelBin link (demo step 6) to show the full config visually.")

# ───────────────────────── 7. Stitching ─────────────────────────
s = new_slide(7)
header(s, "How it stitches together", [[("One click → one trace across three teams", {})]],
       "Real trace from the demo (trace 14911808522202188851): a customer taps Track, and service-c's injected failure is visible end to end.")
code(s, L, 2.25, W, 0.85, [
    [("// injected by the RUM SDK on the browser request", C_)],
    [("traceparent: 00-", CODE_FG), ("0000000000000000cef162d844810c33", H_), ("-1d96e32a32448b12-01      tracestate: dd=s:1;o:rum", CODE_FG)],
], size=13)
lw, tx = 3.7, L + 3.8
tw = R - tx
rows = [("shipping-express · browser", 0, 100, YELLOW, "GET /api/track · 409 ms · 500"),
        ("  └ adt-service-a · jek-my1", 67.3, 26.2, BLUE, "107 ms"),
        ("      └ service-b · jek-my2", 71.7, 17.5, TEAL, "72 ms"),
        ("          └ service-c · jek-cz1", 73.5, 13.8, ORANGE, "57 ms · error")]
for i, (lab, left, width, col, bl) in enumerate(rows):
    y = 3.4 + i * 0.52
    text(s, L, y, lw, 0.42, lab, size=13, font=MONO, anchor=MSO_ANCHOR.MIDDLE)
    box(s, tx, y, tw, 0.42, fill="FFFFFF", fill_alpha=4, radius=0.12)
    box(s, tx + tw * left / 100, y + 0.04, tw * width / 100, 0.34, fill=col, radius=0.15)
    text(s, tx + tw * left / 100 + 0.08, y, max(tw * width / 100, 1.6), 0.42, [[(bl, {"bold": True})]], size=11.5,
         color=BG, anchor=MSO_ANCHOR.MIDDLE)
cw = (W - 0.4) / 3
for i, (q, a) in enumerate([("Where did the time go?", "107 ms in the services; the other ~300 ms was browser queueing and network."),
                            ("Whose error is it?", "service-c, on jek-cz1, with the stack trace attached."),
                            ("What did the customer see?", "Their Session Replay, at the exact moment.")]):
    x = L + i * (cw + 0.2)
    box(s, x, 5.7, cw, 1.2, fill="FFFFFF", fill_alpha=6, line="FFFFFF", line_alpha=14)
    text(s, x + 0.22, 5.8, cw - 0.44, 1.05, [[(q + " ", {"bold": True, "color": WHITE}), (a, {})]], size=14, color="DADCE0", spacing=1.15)
notes(s, "Real trace 14911808522202188851. The RUM SDK injected traceparent; adt-service-a's span is the direct child of the "
         "browser request, and service-c's simulated failure propagates back up. Bar positions use the browser's real "
         "first-byte timing; backend span offsets are approximate.")

# ───────────────────────── 8. Live demo ─────────────────────────
s = new_slide(8)
header(s, "Live demo", [[("Let's follow a shipment", {})]])
demos = [
    ("The customer's view", "Create a shipment, get a quote, track a parcel.", "<MY1_PUBLIC_IP>:8080",
     "http://<MY1_PUBLIC_IP>:8080/"),
    ("Session Replay", "Watch exactly what the customer did and saw.", "RUM → Session Replay",
     "https://app.datadoghq.com/rum/sessions"),
    ("Click → backend trace", "From the RUM request into the A → B → C waterfall.", "RUM → Trace waterfall",
     "https://app.datadoghq.com/rum/sessions"),
    ("Service map", "Dependencies discovered automatically from OTel traces.", "Software Catalog → Map",
     "https://app.datadoghq.com/software"),
    ("Shared dashboard", "Requests, errors, latency, dependencies, hosts and logs on one page.", "Dashboards → ADT Demo",
     "https://app.datadoghq.com/dashboard/lists"),
    ("The Collector config", "The real pipeline, visualised in OTelBin.", "otelbin.io", None),
]
# reuse the exact OTelBin URL from the HTML deck
html = Path(__file__).with_name("otel-demo-slides.html").read_text()
otelbin = html[html.index('href="https://www.otelbin.io/') + 6:]
demos[5] = demos[5][:3] + (otelbin[:otelbin.index('"')],)
cw, ch = (W - 0.4) / 3, 2.4
for i, (t, d, u, href) in enumerate(demos):
    x, y = L + (i % 3) * (cw + 0.2), 1.75 + (i // 3) * (ch + 0.25)
    shp = box(s, x, y, cw, ch, fill="FFFFFF", fill_alpha=6, line="FFFFFF", line_alpha=14)
    shp.click_action.hyperlink.address = href
    text(s, x + 0.25, y + 0.15, 1.0, 0.7, [[(str(i + 1), {"bold": True})]], size=34, color=ORANGE)
    text(s, x + 0.25, y + 0.85, cw - 0.5, 0.45, [[(t, {"bold": True})]], size=19)
    text(s, x + 0.25, y + 1.3, cw - 0.5, 0.6, d, size=13.5, color="DADCE0", spacing=1.12)
    link = text(s, x + 0.25, y + ch - 0.45, cw - 0.5, 0.35, u, size=12, color="B98CFF", font=MONO)
    link.click_action.hyperlink.address = href
notes(s, "Demo order — each card and its link line is clickable:\n" + "\n".join(f"{i + 1}. {d[0]}: {d[3]}" for i, d in enumerate(demos)))

# ───────────────────────── 9. Summary ─────────────────────────
s = new_slide(9)
header(s, "What this means for Acme", [[("Every team sees the ", {}), ("whole journey", {"color": "B98CFF"})]])
cw, ch = (W - 0.4) / 3, 4.15
summary = [("Open by default", ["OpenTelemetry agent and Collector, no proprietary agent on the hosts", "W3C Trace Context end to end",
                                "Portable config, no lock-in at the instrumentation layer"], "Standardise once, across every team"),
           ("Faster triage", ["One trace from browser click to the failing hop", "Logs, host metrics and errors on the same trace ID",
                              "Service map and dependency stats built automatically"], "From war room to the right team, fast"),
           ("Customer-first", ["RUM and Session Replay show what customers experienced", "Errors linked to the backend request that caused them",
                               "Shared dashboards for business and engineering"], "Complaints map straight to code")]
for i, (t, b, o) in enumerate(summary):
    card(s, L + i * (cw + 0.2), 1.75, cw, ch, t, b, ("Benefit", o), hl=True, title_size=21)
box(s, L, 6.15, W, 0.75, fill=ORANGE, fill_alpha=10, radius=0.1)
box(s, L, 6.15, 0.07, 0.75, fill=ORANGE, shape=MSO_SHAPE.RECTANGLE)
text(s, L + 0.25, 6.15, W - 0.4, 0.75, [[("Suggested next step: ", {"bold": True}),
     ("pick one real A → B → C flow and instrument it in a 2-week PoC using this exact pattern.", {})]],
     size=15, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Close: open by default, faster triage, customer-first. Ask for a 2-week PoC on one real A → B → C flow.")

prs.save(OUT)
print("saved", OUT)
